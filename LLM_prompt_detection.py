


# ------------------ SETUP ------------------

# The program works for both PDF and Word. Outputs are available in JSON format (ideal for analysis, APIs, or saving to a file), as well as CSV and Excel.
# The program supports multiple files in one run.

import fitz  # Use PyMuPDF library
import re
import sys
import os
import json
import csv
import openpyxl
from docx import Document # Use Word files

# ------------------ CONFIGURATION ------------------
PROMPT_KEYWORDS = [
    "ignore previous instructions",
    "you are now",
    "disregard earlier",
    "pretend to",
    "output the following",
    "as a language model",
    "repeat after me",
    "print this exactly",
    "simulate response",
    "respond with",
    "complete the prompt"
]

PROMPT_PATTERNS = [re.compile(re.escape(kw), re.IGNORECASE) for kw in PROMPT_KEYWORDS]

# ------------------ PDF SCANNING ------------------
def is_invisible(span):
    color = span.get("color", 0)
    fontsize = span.get("size", 12)
    return color == 16777215 or fontsize < 1

def scan_pdf_visible_text(doc):
    findings = []
    for page_num, page in enumerate(doc):
        blocks = page.get_text("dict")["blocks"]
        for block in blocks:
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    text = span.get("text", "").strip()
                    if not text:
                        continue
                    for pattern in PROMPT_PATTERNS:
                        if pattern.search(text):
                            findings.append({
                                "type": "text",
                                "page": page_num + 1,
                                "text": text,
                                "font_size": span.get("size"),
                                "color": span.get("color"),
                                "invisible": is_invisible(span)
                            })
    return findings

def scan_pdf_metadata(doc):
    findings = []
    metadata = doc.metadata or {}
    for key, value in metadata.items():
        if not value:
            continue
        for pattern in PROMPT_PATTERNS:
            if pattern.search(value):
                findings.append({
                    "type": "metadata",
                    "field": key,
                    "text": value
                })
    return findings

def scan_pdf_javascript(doc):
    findings = []
    for i in range(doc.xref_length()):
        try:
            obj = doc.xref_object(i)
            if "/JavaScript" in obj or "/JS" in obj:
                script_match = re.findall(r"\((.*?)\)", obj, re.DOTALL)
                for script in script_match:
                    for pattern in PROMPT_PATTERNS:
                        if pattern.search(script):
                            findings.append({
                                "type": "javascript",
                                "xref": i,
                                "text": script.strip()
                            })
        except Exception:
            continue
    return findings

def scan_pdf_invisible_and_small_text(doc):
    invisible_texts, small_texts = [], []
    for page_num, page in enumerate(doc):
        blocks = page.get_text("dict")["blocks"]
        for block in blocks:
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    text = span.get("text", "").strip()
                    if not text:
                        continue
                    color = span.get("color", 0)
                    fontsize = span.get("size", 12)
                    if color == 16777215:
                        invisible_texts.append({
                            "page": page_num + 1,
                            "font_size": fontsize,
                            "color": color,
                            "text": text
                        })
                    elif fontsize < 1:
                        small_texts.append({
                            "page": page_num + 1,
                            "font_size": fontsize,
                            "color": color,
                            "text": text
                        })
    return invisible_texts, small_texts

def scan_pdf(path):
    doc = fitz.open(path)
    return {
        "visible_matches": scan_pdf_visible_text(doc),
        "metadata_matches": scan_pdf_metadata(doc),
        "js_matches": scan_pdf_javascript(doc),
        "invisible_texts": scan_pdf_invisible_and_small_text(doc)[0],
        "small_texts": scan_pdf_invisible_and_small_text(doc)[1]
    }

# ------------------ WORD SCANNING ------------------
def scan_docx(path):
    doc = Document(path)
    findings = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        for pattern in PROMPT_PATTERNS:
            if pattern.search(text):
                findings.append({
                    "type": "text",
                    "paragraph": text
                })
    return {"visible_matches": findings}

# ------------------ UNIFIED HANDLER ------------------
def detect_prompts(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return scan_pdf(file_path)
    elif ext == ".docx":
        return scan_docx(file_path)
    else:
        raise ValueError("Unsupported file type")

# ------------------ EXPORT FUNCTIONS ------------------
def save_to_csv(results, out_path):
    with open(out_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["Category", "Details"])
        for key, items in results.items():
            for item in items:
                writer.writerow([key, json.dumps(item, ensure_ascii=False)])

def save_to_excel(results, out_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LLM Prompt Detection"
    ws.append(["Category", "Details"])
    for key, items in results.items():
        for item in items:
            ws.append([key, json.dumps(item, ensure_ascii=False)])
    wb.save(out_path)

def save_consolidated(all_results, csv_out, excel_out):
    # Save CSV
    with open(csv_out, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["File", "Category", "Details"])
        for file, results in all_results.items():
            for key, items in results.items():
                for item in items:
                    writer.writerow([file, key, json.dumps(item, ensure_ascii=False)])

    # Save Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LLM Prompt Detection"
    ws.append(["File", "Category", "Details"])
    for file, results in all_results.items():
        for key, items in results.items():
            for item in items:
                ws.append([file, key, json.dumps(item, ensure_ascii=False)])
    wb.save(excel_out)

# ------------------ CLI ENTRY ------------------
def main(file_paths):
    all_results = {}
    for file_path in file_paths:
        try:
            results = detect_prompts(file_path)
        except Exception as e:
            print(f"❌ Failed to process {file_path}: {e}")
            continue

        all_results[file_path] = results

        # Print JSON results
        print(f"\n=== Results for {file_path} ===")
        print(json.dumps(results, indent=2, ensure_ascii=False))

        # Save per-file CSV & Excel
        base, _ = os.path.splitext(file_path)
        save_to_csv(results, base + "_results.csv")
        save_to_excel(results, base + "_results.xlsx")
        print(f"💾 Results saved to {base}_results.csv and {base}_results.xlsx")

    # Save consolidated report if multiple files
    if len(all_results) > 1:
        save_consolidated(all_results, "consolidated_results.csv", "consolidated_results.xlsx")
        print("\n📊 Consolidated results saved to consolidated_results.csv and consolidated_results.xlsx")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python detect_llm_prompts.py <file1.pdf|file1.docx> [file2.pdf|file2.docx] ...")
        sys.exit(1)
    main(sys.argv[1:])
